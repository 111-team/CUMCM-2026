"""Review checks. Read original deliverables; write only review evidence.

Run after the unmodified run_experiments.py has generated --out /tmp/cumcm_review_rerun.
PYTHONPATH=/tmp/cumcm_review_deps python3 review/verify_review.py
"""
import csv
import json
import sys
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'code'))
from model import load_data, forecast_all, residual_scenarios, optimize, execute

RERUN = Path('/tmp/cumcm_review_rerun')
DT, ETA = 1 / 6, .9
data = load_data(ROOT.parent / 'C题' / '附件')
fc, scores, selected = forecast_all(data)
original = json.loads((ROOT / 'results/summary.json').read_text())
rerun = json.loads((RERUN / 'summary.json').read_text())
out = {'environment': {'numpy': np.__version__}, 'cases': {}, 'workbooks': {}}
all_arrays = {}

for name, meta in original['cases'].items():
    a = dict(np.load(RERUN / (name + '.npz')))
    all_arrays[name] = a
    prices = np.broadcast_to(data['one'][:, 0], (365, 144)) if meta['price_mode'] == 'fixed' else data['price']
    g, ch, ds, em, sp, e = [a[k] for k in ('grid', 'ch', 'dis', 'emergency', 'spill', 'energy')]
    delta = g - a['plan']
    cost = DT * np.sum(prices * (a['plan'] + 1.5 * np.maximum(delta, 0)
                + (.5 if meta['settlement'] == 'additive' else -.5) * np.maximum(-delta, 0)
                + 5 * em), axis=1)
    checks = {
        'total_cost': float(cost[31:].sum()),
        'reported_total_difference': float(cost[31:].sum() - meta['total_cost']),
        'max_power_balance_residual': float(abs(g + data['pv'] + ds + em - data['load'] - ch - sp).max()),
        'max_soc_recursion_residual': float(abs(np.diff(e, axis=1) - ETA * DT * ch + DT * ds / ETA).max()),
        'max_cross_day_gap': float(abs(e[1:, 0] - e[:-1, -1]).max()),
        'min_soc': float(e.min()), 'max_soc': float(e.max()),
        'max_charge': float(ch.max()), 'max_discharge': float(ds.max()),
        'simultaneous_periods': int(np.sum(np.minimum(ch, ds) > 1e-6)),
        'max_daily_cost_residual': float(abs(cost - a['costs'][:, 4]).max()),
    }
    assert abs(checks['reported_total_difference']) < 1e-5
    assert max(checks[k] for k in ('max_power_balance_residual', 'max_soc_recursion_residual', 'max_cross_day_gap', 'max_daily_cost_residual')) < 1e-6
    assert e.min() >= 1200 - 1e-6 and e.max() <= 10800 + 1e-6
    assert ch.max() <= 5000 + 1e-6 and ds.max() <= 5000 + 1e-6
    assert min(g.min(), ch.min(), ds.min(), em.min(), sp.min()) >= -1e-6
    assert checks['simultaneous_periods'] == 0
    out['cases'][name] = checks

for case, filename in [('q2', 'result2.xlsx'), ('q3', 'result3.xlsx'), ('q4_2', 'result4-2.xlsx'), ('q4_3', 'result4-3.xlsx')]:
    a = all_arrays[case]
    wb = load_workbook(ROOT / 'workbooks' / filename, read_only=True, data_only=True)
    checks = {}
    for sheet, key in [('计划购电量', 'plan'), ('调整购电量', 'grid')]:
        if sheet not in wb.sheetnames:
            continue
        rows = list(wb[sheet].values)
        assert [r[0].strftime('%Y-%m-%d') for r in rows[1:]] == data['dates'][31:]
        mat = np.asarray([r[1:] for r in rows[1:]], float)
        checks[sheet + '_max_10min_energy_difference'] = float(abs(mat[:, :144] - a[key][31:] * DT).max())
        checks[sheet + '_max_daily_sum_difference'] = float(abs(mat[:, 144] - mat[:, :144].sum(axis=1)).max())
        expected = a['costs'][31:, 0] if key == 'plan' else a['costs'][31:, :3].sum(axis=1)
        checks[sheet + '_max_cost_difference'] = float(abs(mat[:, 145] - expected).max())
    rows = list(wb['充放电量'].values)[1:]
    assert len(rows) == 334 * 6
    errors = []
    for d in range(334):
        for b in range(6):
            row = rows[d * 6 + b]
            sl = slice(b * 24, (b + 1) * 24)
            errors += [abs(row[2] - DT * a['ch'][d + 31, sl].sum()), abs(row[3] - DT * a['dis'][d + 31, sl].sum())]
            if b < 2:
                errors.append(abs(row[5] - a['energy'][d + 31, 0 if b == 0 else -1]))
    checks['storage_max_difference'] = float(max(errors))
    for row in list(wb['紧急购电量'].values)[1:]:
        if row[0] is not None:
            day = data['dates'].index(row[0].strftime('%Y-%m-%d'))
        if row[1] == '无紧急购电':
            assert abs(a['emergency'][day].sum()) < 1e-6
        else:
            start, end = row[1].split('-')
            idx = [int(h) * 6 + int(m) // 10 for h, m in (start.split(':'), end.split(':'))]
            assert abs(row[2] - DT * a['emergency'][day, idx[0]:idx[1]].sum()) < 1e-6
    assert max(checks.values()) < 1e-6
    out['workbooks'][filename] = checks
    wb.close()

q = dict(np.load(RERUN / 'q1.npz'))
wb = load_workbook(ROOT / 'workbooks/result1.xlsx', read_only=True, data_only=True)
q1_grid = np.asarray([r[1] for r in list(wb['计划购电量'].values)[1:]])
out['q1'] = {'cost': float(DT * data['one'][:, 0] @ q['grid']),
             'grid_kwh': float(q['grid'].sum() * DT),
             'storage_cycle_energy_residual': float(q['dis'].sum() * DT - ETA**2 * q['ch'].sum() * DT),
             'workbook_grid_max_difference': float(abs(q1_grid - q['grid'] * DT).max())}
assert out['q1']['workbook_grid_max_difference'] < 1e-6
for b, row in enumerate(list(wb['充放电量'].values)[1:]):
    sl = slice(b * 24, (b + 1) * 24)
    assert abs(row[1] - q['ch'][sl].sum() * DT) < 1e-6
    assert abs(row[2] - q['dis'][sl].sum() * DT) < 1e-6
wb.close()

with (ROOT / 'results/daily_summary.csv').open(encoding='utf-8-sig') as f:
    rows = list(csv.DictReader(f))
out['csv_max_total_cost_difference'] = max(abs(float(r['total_cost']) - all_arrays[r['case']]['costs'][data['dates'].index(r['date']), 4]) for r in rows)

# Does historical trajectory pairing affect the LP objective? Permute each column independently.
day = data['dates'].index('2025-09-23')
sc = residual_scenarios(data, fc, day, 0, issue=0)
rng = np.random.default_rng(20260911)
shuffled = np.column_stack([rng.permutation(col) for col in sc.T])
initial = all_arrays['q3']['energy'][day, 0]
s1 = optimize(sc, data['one'][:, 0], initial)
s2 = optimize(shuffled, data['one'][:, 0], initial)
out['scenario_column_permutation'] = {'date': data['dates'][day], 'original_objective': s1['objective'], 'permuted_objective': s2['objective'], 'objective_difference': s2['objective'] - s1['objective'], 'max_grid_difference': float(abs(s2['grid'] - s1['grid']).max())}
for label, sol in [('original', s1), ('permuted', s2)]:
    feedback = execute(sol['grid'], data['load'][day], data['pv'][day], initial)
    out['scenario_column_permutation'][label + '_fixed_grid_day_cost'] = float(DT * data['one'][:, 0] @ (sol['grid'] + 5 * feedback['emergency']))

# Small, explicitly synthetic counterexample for greedy feedback; not contest data.
ex = execute(np.zeros(2), np.full(2, 6.), np.zeros(2), 1200 + 1 / ETA)
prices = np.array([.4, 1.4])
out['synthetic_feedback_counterexample'] = {'greedy_cost': float(5 * DT * prices @ ex['emergency']), 'feasible_reserve_first_period_cost': 2., 'greedy_emergency_kwh': (DT * ex['emergency']).tolist(), 'initial_soc': 1200 + 1 / ETA}

# Compare original zero-hour and updated forecasts over identical future six-hour blocks.
out['pv_forecast_common_horizons'] = {}
out['pv_forecast_full_day_mae_kw'] = {
    'historical': float(abs(fc['pv'][31:] - data['pv'][31:]).mean()),
    'zero_issue': float(abs(fc['issued'][31:, 0] - data['pv'][31:]).mean()),
}
for issue in range(4):
    sl = slice(issue * 36, (issue + 1) * 36)
    truth = data['pv'][31:, sl]
    out['pv_forecast_common_horizons'][str(issue * 6)] = {
        'historical_mae_kw': float(abs(fc['pv'][31:, sl] - truth).mean()),
        'zero_issue_mae_kw': float(abs(fc['issued'][31:, 0, sl] - truth).mean()),
        'current_issue_mae_kw': float(abs(fc['issued'][31:, issue, sl] - truth).mean()),
    }

delta = all_arrays['q3_daily']['costs'][31:, 4] - all_arrays['q3']['costs'][31:, 4]
out['update_cost_by_month'] = {str(m): float(delta[[int(d[5:7]) == m for d in data['dates'][31:]]].sum()) for m in range(2, 13)}
out['update_days'] = {'cheaper': int((delta > 1e-6).sum()), 'more_expensive': int((delta < -1e-6).sum()), 'same': int((abs(delta) <= 1e-6).sum())}

# Reproduce the paper's future perturbation check; it is a forecast check only.
changed = {k: v.copy() if isinstance(v, np.ndarray) else v for k, v in data.items()}
for k in ('load', 'pv', 'price'):
    changed[k][180:] *= 1.5
changed['pv_hour'][180:] *= .5
fc2, _, selected2 = forecast_all(changed)
assert selected == selected2
assert all(np.array_equal(fc[k][:180], fc2[k][:180]) for k in fc)
out['forecast_future_perturbation'] = 'passed; this does not by itself test decisions'

target = ROOT / 'review/verification_review.json'
target.write_text(json.dumps(out, ensure_ascii=False, indent=2))
print(json.dumps(out, ensure_ascii=False, indent=2))
