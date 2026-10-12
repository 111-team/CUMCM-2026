"""按附件5填充5份主结果和2份价格已知补充表；全天购电费包含紧急费。"""
import argparse,json,csv
from pathlib import Path
from copy import copy
from datetime import datetime
import numpy as np
from openpyxl import load_workbook
from openpyxl.styles import Alignment,Font
from model import DT,EINIT


def clock(t):
    minutes=t*10
    return f'{minutes//60}:{minutes%60:02d}'


def interval(t):
    return clock(t)+'-'+clock(t+1)


def events(power):
    out=[];start=None
    for t in range(145):
        active=t<144 and power[t]>1e-6
        if active and start is None:start=t
        if not active and start is not None:
            out.append((clock(start)+'-'+clock(t),float(power[start:t].sum()*DT)))
            start=None
    return out


def reset_body(ws):
    for merged in list(ws.merged_cells.ranges):ws.unmerge_cells(str(merged))
    styles=[copy(ws.cell(2,c)._style) for c in range(1,ws.max_column+1)]
    if ws.max_row>1:ws.delete_rows(2,ws.max_row-1)
    return styles


def append_row(ws,values,styles):
    ws.append(values)
    r=ws.max_row
    for c,v in enumerate(values,1):
        cell=ws.cell(r,c)
        if c<=len(styles):cell._style=copy(styles[c-1])
        if isinstance(v,(float,np.floating)):cell.number_format='0.0000'
        if isinstance(v,datetime):cell.number_format='yyyy/m/d'
        cell.alignment=Alignment(horizontal='center',vertical='center')
    ws.row_dimensions[r].height=20


def main():
    parser=argparse.ArgumentParser();parser.add_argument('--data',type=Path,required=True)
    parser.add_argument('--root',type=Path,default=Path(__file__).resolve().parents[1]);args=parser.parse_args()
    root=args.root;result=root/'results';dest=root/'结果';dest.mkdir(exist_ok=True)
    meta=json.loads((result/'summary.json').read_text(encoding='utf-8'));dates=meta['dates']
    q1=np.load(result/'q1.npz')
    wb=load_workbook(args.data/'附件5'/'result1.xlsx')
    ws=wb['计划购电量'];styles=reset_body(ws)
    for t in range(144):append_row(ws,[interval(t),float(q1['grid'][t]*DT)],styles)
    ws=wb['充放电量'];styles=reset_body(ws)
    for block in range(6):
        sl=slice(block*24,(block+1)*24)
        append_row(ws,[f'{block*4}:00-{(block+1)*4}:00',float(q1['ch'][sl].sum()*DT),float(q1['dis'][sl].sum()*DT),('0:00' if block==0 else '24:00' if block==1 else None),(EINIT if block==0 else float(q1['energy'][-1]) if block==1 else None)],styles)
    wb.save(dest/'result1.xlsx')
    daily=[];checks={}
    cases=[('q2','result2.xlsx'),('q3','result3.xlsx'),('q4_2','result4-2.xlsx'),('q4_3','result4-3.xlsx'),
           ('q4_2_known','result4-2.xlsx'),('q4_3_known','result4-3.xlsx')]
    supplement=dest/'补充';supplement.mkdir(exist_ok=True)
    for case,filename in cases:
        output=(supplement/('补充_价格已知_'+filename)) if case.endswith('_known') else dest/filename
        a=np.load(result/(case+'.npz'));wb=load_workbook(args.data/'附件5'/filename)
        for sheet,key in [('计划购电量','plan')]+([('调整购电量','grid')] if '调整购电量' in wb.sheetnames else []):
            ws=wb[sheet];styles=reset_body(ws)
            for t in range(144):ws.cell(1,t+2,interval(t))
            for d in range(31,365):
                # 两个购电工作表显示同一当日总费用，不能把两列再次相加。
                cost=a['costs'][d,4]
                append_row(ws,[datetime.fromisoformat(dates[d])]+(a[key][d]*DT).tolist()+[float(a[key][d].sum()*DT),float(cost)],styles)
            ws.freeze_panes='B2'
        ws=wb['充放电量'];styles=reset_body(ws)
        for d in range(31,365):
            for block in range(6):
                sl=slice(block*24,(block+1)*24)
                append_row(ws,[datetime.fromisoformat(dates[d]) if block==0 else None,f'{block*4}:00-{(block+1)*4}:00',float(a['ch'][d,sl].sum()*DT),float(a['dis'][d,sl].sum()*DT),('0:00' if block==0 else '24:00' if block==1 else None),(float(a['energy'][d,0]) if block==0 else float(a['energy'][d,-1]) if block==1 else None)],styles)
        ws.freeze_panes='C2'
        ws=wb['紧急购电量'];styles=reset_body(ws)
        for d in range(31,365):
            for i,(times,kwh) in enumerate(events(a['emergency'][d]) or [('无紧急购电',0.)]):
                append_row(ws,[datetime.fromisoformat(dates[d]) if i==0 else None,times,kwh],styles)
            daily.append([case,dates[d],*a['costs'][d].tolist(),float(a['plan'][d].sum()*DT),float(a['grid'][d].sum()*DT),float(a['emergency'][d].sum()*DT),float(a['energy'][d,0]),float(a['energy'][d,-1])])
        for ws in wb:
            ws.column_dimensions['A'].width=max(ws.column_dimensions['A'].width or 12,16)
            if ws.title=='紧急购电量':ws.column_dimensions['B'].width=24;ws.column_dimensions['C'].width=18
            if ws.title=='充放电量':
                for c in 'BCDEF':ws.column_dimensions[c].width=17
        wb.save(output)
        # 独立重新读取保存文件，逐日核对电量总和及费用列。
        check=load_workbook(output,read_only=True,data_only=True)
        for sheet,key in [('计划购电量','plan')]+([('调整购电量','grid')] if '调整购电量' in check.sheetnames else []):
            rows=list(check[sheet].values)
            assert list(rows[0][1:145])==[interval(t) for t in range(144)]
            assert [v[0].strftime('%Y-%m-%d') for v in rows[1:]]==dates[31:]
            matrix=np.asarray([row[1:] for row in rows[1:]],float)
            expected_cost=a['costs'][31:,4]
            assert np.max(abs(matrix[:,:144]-a[key][31:]*DT))<1e-6
            assert np.max(abs(matrix[:,:144].sum(axis=1)-matrix[:,144]))<1e-6
            assert np.max(abs(matrix[:,145]-expected_cost))<1e-6
        rows=list(check['充放电量'].values)[1:]
        for col,key in [(2,'ch'),(3,'dis')]:
            expected=a[key][31:].reshape(334,6,24).sum(axis=2).ravel()*DT
            assert np.max(abs(np.array([r[col] for r in rows])-expected))<1e-6
        assert np.max(abs(np.array([rows[d*6][5] for d in range(334)])-a['energy'][31:,0]))<1e-6
        assert np.max(abs(np.array([rows[d*6+1][5] for d in range(334)])-a['energy'][31:,-1]))<1e-6
        expected_events=[(times,kwh) for d in range(31,365) for times,kwh in (events(a['emergency'][d]) or [('无紧急购电',0.)])]
        stored_events=list(check['紧急购电量'].values)[1:]
        assert len(expected_events)==len(stored_events)
        for stored,(times,kwh) in zip(stored_events,expected_events):
            assert stored[1]==times and abs(stored[2]-kwh)<1e-6
        checks[str(output.relative_to(dest))]={'days':334,'quantity_cost_storage_events':'passed',
            'settlement':meta['cases'][case]['settlement'],'case':case,'total_cost':float(a['costs'][31:,4].sum())}
        check.close()
    with (result/'daily_summary.csv').open('w',newline='',encoding='utf-8-sig') as f:
        writer=csv.writer(f);writer.writerow(['case','date','plan_cost','up_cost','down_cost','emergency_cost','total_cost','plan_kwh','grid_kwh','emergency_kwh','start_energy','end_energy']);writer.writerows(daily)
    check=load_workbook(dest/'result1.xlsx',read_only=True,data_only=True)
    rows=list(check['计划购电量'].values)[1:]
    assert [r[0] for r in rows]==[interval(t) for t in range(144)]
    assert np.max(abs(np.array([r[1] for r in rows])-q1['grid']*DT))<1e-6
    rows=list(check['充放电量'].values)[1:]
    for col,key in [(1,'ch'),(2,'dis')]:
        assert np.max(abs(np.array([r[col] for r in rows])-q1[key].reshape(6,24).sum(axis=1)*DT))<1e-6
    assert abs(rows[0][4]-EINIT)<1e-6 and abs(rows[1][4]-q1['energy'][-1])<1e-6
    check.close();checks['result1.xlsx']={'quantity_storage':'passed'}
    (result/'workbook_verification.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2),encoding='utf-8')
    print('Exported and reloaded 5 main and 2 supplementary workbooks under 结果/; quantities, total costs, storage and events passed.')

if __name__=='__main__':main()
